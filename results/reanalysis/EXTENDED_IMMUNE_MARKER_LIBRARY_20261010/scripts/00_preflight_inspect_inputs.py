#!/usr/bin/env python3
"""Read-only preflight inspection for the extended immune marker audit."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import anndata as ad
import pandas as pd
from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser.parse_args()


def sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def workbook_summary(path: Path) -> tuple[dict[str, Any], dict[str, pd.DataFrame]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    recovered: dict[str, pd.DataFrame] = {}
    for worksheet in workbook.worksheets:
        reported_dimension = worksheet.calculate_dimension()
        if reported_dimension == "A1:A1":
            # This source workbook contains a stale <dimension ref="A1"/>
            # despite thousands of valid cells. Reset only the in-memory
            # reader bounds; the original XLSX is never modified.
            worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
        preview = [
            [json_value(value) for value in row[:20]]
            for row in rows[:15]
        ]
        headers = [str(value) if value is not None else "" for value in rows[0]]
        frame = pd.DataFrame(rows[1:], columns=headers)
        recovered[worksheet.title] = frame
        sheets.append(
            {
                "title": worksheet.title,
                "reported_dimension": reported_dimension,
                "recovered_row_count": len(frame),
                "recovered_column_count": len(frame.columns),
                "columns": list(map(str, frame.columns)),
                "level_counts": (
                    frame["level"].astype(str).value_counts(dropna=False).to_dict()
                    if "level" in frame.columns else {}
                ),
                "celltype_counts": (
                    frame.groupby(["level", "celltype"], dropna=False).size()
                    .reset_index(name="record_n").astype(str).to_dict("records")
                    if {"level", "celltype"}.issubset(frame.columns) else []
                ),
                "preview_first_15_rows_first_20_columns": preview,
            }
        )
    workbook.close()
    return {"sheet_count": len(sheets), "sheets": sheets}, recovered


def h5ad_summary(path: Path) -> dict[str, Any]:
    data = ad.read_h5ad(path, backed="r")
    try:
        obs = data.obs
        var = data.var
        return {
            "shape": [int(data.n_obs), int(data.n_vars)],
            "obs_columns": list(map(str, obs.columns)),
            "var_columns": list(map(str, var.columns)),
            "obs_head": obs.head(5).reset_index().astype(str).to_dict("records"),
            "var_head": var.head(5).reset_index().astype(str).to_dict("records"),
            "obs_names_unique": bool(data.obs_names.is_unique),
            "var_names_unique": bool(data.var_names.is_unique),
            "x_storage": type(data.X).__name__,
        }
    finally:
        data.file.close()


def manifest_row(path: Path, purpose: str) -> dict[str, Any]:
    stat = path.stat()
    return {
        "absolute_path": str(path.resolve()),
        "file_name": path.name,
        "size_bytes": stat.st_size,
        "modified_time": pd.Timestamp(stat.st_mtime, unit="s", tz="UTC").isoformat(),
        "sha256": sha256(path),
        "analysis_purpose": purpose,
    }


def main() -> None:
    args = parse_args()
    root = args.project_root.resolve()
    output = args.output_dir.resolve()
    qc_dir = output / "qc"
    qc_dir.mkdir(parents=True, exist_ok=True)

    atlas_candidates = [
        root / "results" / "reanalysis" / "bin50_input"
        / "BI_PF_ILD_atlas_marker_gene_table_260709213053.xlsx",
        root / "results" / "visium way" / "06_hvg_audit"
        / "BI_PF_ILD_atlas_marker_gene_table_260709213053.xlsx",
    ]
    atlas_paths = [path for path in atlas_candidates if path.exists()]
    secondary = list(root.rglob("BI_PF_ILD_atlas_marker_gene_table_260723131417.xlsx"))
    raw_h5ad = (
        root / "results" / "reanalysis" / "bin50_input"
        / "BIN50_joint_reanalysis_tissue_raw_counts_common_genes.h5ad"
    )
    gene_map = (
        root / "results" / "reanalysis" / "bin50_representation_benchmark"
        / "ensembl_gene_symbol_mapping.csv"
    )
    legacy_marker_source = root / "scripts" / "1049_prepare_immune_subniche_validation.R"
    required = atlas_paths + [raw_h5ad, gene_map, legacy_marker_source]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required inputs: " + "; ".join(missing))

    atlas_summaries = {}
    recovered_primary: dict[str, pd.DataFrame] = {}
    for index, path in enumerate(atlas_paths):
        summary, recovered = workbook_summary(path)
        atlas_summaries[str(path)] = summary
        if index == 0:
            recovered_primary = recovered
    hashes = {str(path): sha256(path) for path in atlas_paths}
    report = {
        "project_root": str(root),
        "output_dir": str(output),
        "primary_atlas_candidates": [str(path) for path in atlas_paths],
        "primary_atlas_hashes": hashes,
        "primary_copies_identical": len(set(hashes.values())) == 1,
        "secondary_atlas_candidates": [str(path) for path in secondary],
        "atlas_workbooks": atlas_summaries,
        "raw_h5ad": {"path": str(raw_h5ad), **h5ad_summary(raw_h5ad)},
        "gene_map": {
            "path": str(gene_map),
            "columns": list(pd.read_csv(gene_map, nrows=5).columns),
            "head": pd.read_csv(gene_map, nrows=5, dtype=str).fillna("").to_dict("records"),
        },
        "legacy_small_marker_panel_source": {
            "path": str(legacy_marker_source),
            "sha256": sha256(legacy_marker_source),
        },
    }
    (qc_dir / "00_PREFLIGHT_INPUT_INSPECTION.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for sheet_name, frame in recovered_primary.items():
        safe_sheet = "".join(
            character if character.isalnum() else "_" for character in sheet_name
        ).strip("_") or "sheet"
        frame.to_csv(
            qc_dir / f"00_ATLAS_CELL_DATA_RECOVERED_{safe_sheet}.tsv",
            sep="\t",
            index=False,
        )

    manifest = []
    for path in atlas_paths:
        manifest.append(manifest_row(path, "Primary atlas marker input; read-only copy"))
    for path in secondary:
        manifest.append(manifest_row(path, "Optional stromal-background reference only"))
    manifest.append(manifest_row(raw_h5ad, "QCed 21-patient sparse Bin50 raw-count matrix"))
    manifest.append(manifest_row(gene_map, "Ensembl-to-gene-symbol mapping"))
    manifest.append(manifest_row(
        legacy_marker_source,
        "Source code defining the previously used small immune marker panels",
    ))
    pd.DataFrame(manifest).to_csv(output / "input_manifest.tsv", sep="\t", index=False)

    print(json.dumps({
        "atlas_copy_n": len(atlas_paths),
        "atlas_copies_identical": report["primary_copies_identical"],
        "secondary_atlas_n": len(secondary),
        "atlas_sheet_names": {
            str(path): [sheet["title"] for sheet in summary["sheets"]]
            for path, summary in [(Path(key), value) for key, value in atlas_summaries.items()]
        },
        "raw_shape": report["raw_h5ad"]["shape"],
        "obs_columns": report["raw_h5ad"]["obs_columns"],
        "var_columns": report["raw_h5ad"]["var_columns"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
